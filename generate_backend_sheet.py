import re
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def main():
    with open('index.html', 'r', encoding='utf-8') as f:
        content = f.read()

    def extract_array(var_name):
        pattern = rf'let\s+{var_name}\s*=\s*(\[[\s\S]*?\]);'
        m = re.search(pattern, content)
        if not m:
            raise ValueError(f"Could not find {var_name} in index.html")
        return json.loads(m.group(1))

    schedule_data = extract_array('SCHEDULE_DATA')
    territory_data = extract_array('TERRITORY_DATA')
    meetings_data = extract_array('MEETINGS_DATA')
    hospitality_data = extract_array('HOSPITALITY_DATA')

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='265343', end_color='265343', fill_type='solid') # Brand Pine Green
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')
    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )

    datasets = [
        ('Schedule', schedule_data, ['dayKey', 'dayLabel', 'fullDate', 'time', 'role', 'activity', 'companion', 'status', 'note']),
        ('Territory', territory_data, ['day', 'time', 'territory', 'meetingPlace', 'guide', 'notes']),
        ('Meetings', meetings_data, ['title', 'dayTime', 'place', 'attendees', 'desc']),
        ('Hospitality', hospitality_data, ['day', 'date', 'breakfast', 'lunch', 'dinner'])
    ]

    for title, rows, headers in datasets:
        ws = wb.create_sheet(title=title)
        ws.views.sheetView[0].showGridLines = True

        # Header row
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border
        
        ws.row_dimensions[1].height = 26

        # Data rows
        for row_idx, row_data in enumerate(rows, 2):
            ws.row_dimensions[row_idx].height = 20
            for col_idx, header in enumerate(headers, 1):
                val = row_data.get(header, '')
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = Font(name='Calibri', size=10)
                cell.alignment = left_align
                cell.border = thin_border

        # Adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output_filename = 'CO_Visit_Backend_Data.xlsx'
    wb.save(output_filename)
    print(f"Successfully generated {output_filename} with {len(schedule_data)} schedule slots, {len(territory_data)} territories, {len(meetings_data)} meetings, {len(hospitality_data)} hospitality rows.")

if __name__ == '__main__':
    main()
